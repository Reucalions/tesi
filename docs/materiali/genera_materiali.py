"""Impagina i materiali di presentazione; nessuna inferenza o modifica del prototipo.

Usare l'ambiente separato descritto nel README di questa cartella.
Slide testuali e diagrammi nel PowerPoint rimangono modificabili.
"""

import html
import json
import math
from functools import partial
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocInches
from docx.shared import Pt as DocPt
from docx.shared import RGBColor as DocRGBColor
from lxml import etree
from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    XPreformatted,
)

HERE = Path(__file__).resolve().parent
# PT Sans e ottanio: ispirazione dal manuale del marchio Poliba (README).
FONT_FAMILY = "PT Sans"
INK, TEAL, LIGHT, MUTED, WHITE = "12263A", "008195", "EAF4F5", "526577", "FFFFFF"
W, H = 13.333, 7.5
SLIDES = json.loads((HERE / "slides.json").read_text())


def ppt_text(slide, text, x, y, w, h, size=22, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = FONT_FAMILY
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(5)
    return shape


def box(slide, x, y, w, h, fill=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    return shape


def arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = RGBColor.from_string(TEAL)
    line.line.width = Pt(2)
    line._element.spPr.get_or_add_ln().append(OxmlElement("a:tailEnd"))
    line._element.spPr.ln[-1].set("type", "triangle")


def dag_elements(rect, text, line):
    """Un solo insieme di coordinate per diagramma PPTX e figura della relazione."""
    text("Input: StaticTelemetry", 0.6, 1.7, 3.5, 0.3, 15, MUTED, False)
    rect(0.6, 2.6, 3.4, 0.85, LIGHT)
    rect(4.97, 2.6, 3.4, 0.85, LIGHT)
    rect(9.33, 2.6, 3.4, 0.85, LIGHT)
    rect(4.77, 5.2, 3.8, 0.85, TEAL)
    for label, x in [
        ("TelemetryAgent", 0.8),
        ("VulnerabilityAgent", 5.17),
        ("CountermeasureAgent", 9.53),
    ]:
        text(label, x, 2.88, 3.0, 0.4, 19, INK, True)
    text("StrategicAgent", 5.36, 5.48, 3.0, 0.4, 22, WHITE, True)
    line(4.0, 3.02, 4.97, 3.02)
    line(8.37, 3.02, 9.33, 3.02)
    # La dipendenza diretta RuntimeEvidence -> Countermeasure passa sopra i nodi.
    line(2.3, 2.6, 2.3, 2.25)
    line(2.3, 2.25, 11.03, 2.25)
    line(11.03, 2.25, 11.03, 2.6)
    text("RuntimeEvidence per Countermeasure", 4.5, 1.83, 5.8, 0.35, 16, TEAL, False)
    text("lookup", 4.05, 3.22, 0.9, 0.35, 12, MUTED, False)
    text("report", 8.43, 3.22, 0.85, 0.35, 12, MUTED, False)
    line(2.3, 3.45, 4.77, 5.53)
    line(6.67, 3.45, 6.67, 5.2)
    line(11.03, 3.45, 8.57, 5.53)
    text("RuntimeEvidence", 0.8, 4.25, 3.0, 0.4, 18, TEAL, False)
    text("VulnerabilityReport", 6.9, 4.0, 3.5, 0.4, 17, TEAL, False)
    text("CandidateStrategies", 9.45, 4.65, 3.4, 0.4, 17, TEAL, False)


def create_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Prototipo multi-agente per Agentic CTI"
    prs.core_properties.subject = "Mockup CAMEL/Ollama: presentazione da 15 minuti"
    for i, data in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        box(slide, 0.0, 0.0, W, 0.10, TEAL)
        ppt_text(slide, "AGENTIC CTI  /  PROTOTIPO DI TESI", 0.6, 0.3, 11, 0.25, 11, TEAL, True)
        ppt_text(slide, data["title"], 0.6, 0.78, 12.1, 1.0, 30, INK, True)
        ppt_text(
            slide,
            data["subtitle"],
            0.6,
            1.78 if data["kind"] != "dag" else 1.36,
            12.1,
            0.55,
            18,
            MUTED,
        )
        kind = data["kind"]
        if kind == "dag":
            dag_elements(partial(box, slide), partial(ppt_text, slide), partial(arrow, slide))
        elif kind == "table":
            table = slide.shapes.add_table(
                len(data["rows"]) + 1,
                len(data["headers"]),
                Inches(0.6),
                Inches(2.6),
                Inches(12.1),
                Inches(3.3),
            ).table
            for row_id, row in enumerate([data["headers"], *data["rows"]]):
                for col_id, value in enumerate(row):
                    cell = table.cell(row_id, col_id)
                    cell.text = value
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(TEAL if row_id == 0 else LIGHT)
                    cell.margin_left = Inches(0.15)
                    cell.margin_top = Inches(0.12)
                    for p in cell.text_frame.paragraphs:
                        p.font.name = FONT_FAMILY
                        p.font.size = Pt(18)
                        p.font.bold = row_id == 0
                        p.font.color.rgb = RGBColor.from_string(WHITE if row_id == 0 else INK)
        else:
            y = 2.75
            for n, (heading, body) in enumerate(data["cards"]):
                x = 0.6 + n * 4.1
                box(slide, x, y, 3.9, 2.25)
                ppt_text(slide, heading, x + 0.23, y + 0.3, 3.42, 0.6, 23, TEAL, True)
                ppt_text(slide, body, x + 0.23, y + 1.0, 3.42, 1.12, 21)
            if kind == "demo":
                box(slide, 0.6, 5.35, 12.1, 1.0, INK)
                ppt_text(slide, data["command"], 0.8, 5.55, 11.7, 0.75, 15, WHITE)
        if data.get("callout"):
            ppt_text(
                slide,
                data["callout"],
                0.65,
                6.12 if kind == "table" else 5.55,
                12.0,
                0.78,
                20,
                TEAL,
                True,
            )
        ppt_text(slide, data["footer"], 0.6, 7.02, 11.7, 0.28, 11, MUTED)
        ppt_text(slide, f"{i:02d}", 12.3, 7.02, 0.5, 0.28, 12, TEAL, True)
        timing = (
            f"Tempo indicativo: {data['seconds']} secondi." if data["seconds"] else "Appendice."
        )
        slide.notes_slide.notes_text_frame.text = timing + "\n\n" + data["notes"]
        for paragraph in slide.notes_slide.notes_text_frame.paragraphs:
            paragraph.font.name = FONT_FAMILY
    set_theme_fonts(prs.part.package)
    prs.save(HERE / "presentazione.pptx")


def set_theme_fonts(package):
    """Allinea anche il tema Office: il nuovo testo eredita la famiglia scelta."""
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for part in package.iter_parts():
        if part.content_type.endswith("theme+xml"):
            root = etree.fromstring(part.blob)
            for element in root.xpath("//a:fontScheme/*/a:latin", namespaces=namespace):
                element.set("typeface", FONT_FAMILY)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8")


def register_fonts():
    """Incorpora nei PDF i font distribuiti col progetto, senza fallback di sistema."""
    for suffix, variant in [
        ("", "Regular"),
        ("Bold", "Bold"),
        ("Italic", "Italic"),
        ("BoldItalic", "BoldItalic"),
    ]:
        path = HERE / "fonts" / f"PT_Sans-Web-{variant}.ttf"
        pdfmetrics.registerFont(TTFont("Body" + suffix, str(path)))
    pdfmetrics.registerFontFamily(
        "Body", normal="Body", bold="BodyBold", italic="BodyItalic", boldItalic="BodyBoldItalic"
    )
    return "Body"


def create_figure(font):
    c = canvas.Canvas(str(HERE / "architettura.pdf"), pagesize=(960, 432))
    scale = 72

    def rect(x, y, w, h, fill):
        c.setFillColor(colors.HexColor("#" + fill))
        c.roundRect(x * scale, (6.7 - y - h) * scale, w * scale, h * scale, 8, fill=1, stroke=0)

    def text(value, x, y, w, h, size, color, bold):
        c.setFillColor(colors.HexColor("#" + color))
        c.setFont("BodyBold" if bold and font == "Body" else font, size)
        c.drawString(x * scale, (6.7 - y) * scale - size, value)

    def line(x1, y1, x2, y2):
        ax, ay, bx, by = x1 * scale, (6.7 - y1) * scale, x2 * scale, (6.7 - y2) * scale
        c.setStrokeColor(colors.HexColor("#" + TEAL))
        c.setFillColor(colors.HexColor("#" + TEAL))
        c.setLineWidth(2)
        c.line(ax, ay, bx, by)
        angle = math.atan2(by - ay, bx - ax)
        p = c.beginPath()
        p.moveTo(bx, by)
        p.lineTo(bx - 7 * math.cos(angle - 0.4), by - 7 * math.sin(angle - 0.4))
        p.lineTo(bx - 7 * math.cos(angle + 0.4), by - 7 * math.sin(angle + 0.4))
        p.close()
        c.drawPath(p, fill=1, stroke=0)

    dag_elements(rect, text, line)
    c.save()
    doc = pymupdf.open(HERE / "architettura.pdf")
    doc[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(HERE / "architettura.png")
    doc.close()


def create_report(font):
    source = (HERE.parent / "relazione-mockup.md").read_text()
    tokens = MarkdownIt("commonmark").enable("table").parse(source)
    assert tokens[0].type == "heading_open" and tokens[0].tag == "h1"
    title = tokens[1].content
    tokens = tokens[3:]  # Il titolo principale viene collocato sul frontespizio.
    cover = json.loads((HERE / "frontespizio.json").read_text())
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    styles["Normal"].fontSize = 10
    styles["Normal"].leading = 14
    styles["Normal"].spaceAfter = 7
    styles["Normal"].alignment = TA_LEFT
    styles["Normal"].allowWidows = 0
    for name, size in [("Title", 24), ("Heading1", 17), ("Heading2", 13), ("Heading3", 11)]:
        styles[name].fontName = "BodyBold"
        styles[name].fontSize = size
        styles[name].leading = size + 4
        styles[name].spaceBefore = 12
        styles[name].spaceAfter = 7
        styles[name].keepWithNext = True
        styles[name].textColor = colors.HexColor("#" + TEAL)
    code_style = ParagraphStyle(
        "code",
        fontName="Courier",
        fontSize=7.2,
        leading=10,
        backColor=colors.HexColor("#F0F3F5"),
        spaceAfter=10,
    )
    cell_style = ParagraphStyle(
        "cell", parent=styles["Normal"], fontSize=8, leading=11, spaceAfter=0
    )
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = DocPt(A4[0]), DocPt(A4[1])
    section.left_margin = section.right_margin = DocPt(60)
    section.top_margin = section.bottom_margin = DocPt(48)
    document.core_properties.title = title
    document.core_properties.author = cover["autore"]
    document.core_properties.subject = cover["tipo"]
    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = FONT_FAMILY
    document.styles["Normal"].font.size = DocPt(10)
    for name, size in [("Title", 24), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
        style = document.styles[name]
        style.font.size = DocPt(size)
        style.font.bold = True
        style.font.color.rgb = DocRGBColor.from_string(TEAL)
    logo = HERE / "assets" / "logo-poliba.png"
    logo_width = 300
    logo_image = Image(str(logo), width=logo_width)
    logo_image.drawHeight = logo_width * logo_image.imageHeight / logo_image.imageWidth
    cover_style = ParagraphStyle(
        "cover", parent=styles["Normal"], alignment=TA_CENTER, fontSize=12, leading=18
    )
    cover_title = ParagraphStyle(
        "cover-title",
        parent=cover_style,
        fontName="BodyBold",
        fontSize=25,
        leading=32,
        textColor=colors.HexColor("#" + TEAL),
    )
    story = [Spacer(1, 28), logo_image, Spacer(1, 32)]

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = DocPt(28)
    p.paragraph_format.space_after = DocPt(32)
    p.add_run().add_picture(str(logo), width=DocPt(logo_width))

    def cover_line(text, *, size=12, bold=False, before=0, after=7):
        style = cover_title if size == 25 else cover_style
        markup = html.escape(text)
        if bold:
            markup = "<b>" + markup + "</b>"
        story.extend([Spacer(1, before), Paragraph(markup, style), Spacer(1, after)])
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = DocPt(before)
        paragraph.paragraph_format.space_after = DocPt(after)
        run = paragraph.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.size = DocPt(size)
        run.bold = bold
        if size == 25:
            run.font.color.rgb = DocRGBColor.from_string(TEAL)

    cover_line(cover["ateneo"], bold=True)
    cover_line(cover["corso"])
    cover_line(cover["tipo"], before=35)
    cover_line(title, size=25, bold=True, before=12, after=12)
    cover_line(cover["sottotitolo"])
    cover_line("Autore: " + cover["autore"], before=48)
    cover_line("Relatore: " + cover["relatore"])
    cover_line(cover["data"], before=35)
    story.append(PageBreak())
    document.add_page_break()

    def inline(token):
        parts = []
        links = []
        for child in token.children or []:
            if child.type == "text":
                parts.append(html.escape(child.content))
            elif child.type == "code_inline":
                parts.append('<font name="Courier">' + html.escape(child.content) + "</font>")
            elif child.type in ("strong_open", "strong_close"):
                parts.append("<b>" if child.type.endswith("open") else "</b>")
            elif child.type in ("em_open", "em_close"):
                parts.append("<i>" if child.type.endswith("open") else "</i>")
            elif child.type in ("softbreak", "hardbreak"):
                parts.append(" " if child.type == "softbreak" else "<br/>")
            elif child.type == "link_open":
                href = child.attrGet("href") or ""
                external = href.startswith(("https://", "http://"))
                links.append(external)
                if external:
                    parts.append('<link href="' + html.escape(href, quote=True) + '">')
            elif child.type == "link_close" and links.pop():
                parts.append("</link>")
        return "".join(parts)

    def word_inline(paragraph, token, prefix=""):
        """Conserva enfasi e codice inline anche nella copia modificabile."""
        if prefix:
            paragraph.add_run(prefix)
        bold, italic = False, False
        for child in token.children or []:
            if child.type in ("strong_open", "strong_close"):
                bold = child.type.endswith("open")
            elif child.type in ("em_open", "em_close"):
                italic = child.type.endswith("open")
            elif child.type in ("text", "code_inline", "softbreak", "hardbreak"):
                value = " " if child.type == "softbreak" else child.content
                value = "\n" if child.type == "hardbreak" else value
                run = paragraph.add_run(value)
                run.bold, run.italic = bold, italic
                if child.type == "code_inline":
                    run.font.name = "Courier New"
                    run.font.size = DocPt(8)

    i, lists = 0, []
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            content = tokens[i + 1]
            level = int(token.tag[1:])
            style = "Title" if level == 1 else f"Heading{level - 1}"
            story.append(Paragraph(inline(content), styles[style]))
            document.add_heading(content.content, level=level - 1)
            i += 3
            continue
        if token.type == "table_open":
            rows, row = [], []
            i += 1
            while tokens[i].type != "table_close":
                if tokens[i].type == "tr_open":
                    row = []
                elif tokens[i].type == "inline":
                    row.append(tokens[i])
                elif tokens[i].type == "tr_close":
                    rows.append(row)
                i += 1
            widths = [475 / len(rows[0])] * len(rows[0])
            table = Table(
                [
                    [
                        Paragraph(
                            "<b>" + inline(t) + "</b>" if row_id == 0 else inline(t), cell_style
                        )
                        for t in row
                    ]
                    for row_id, row in enumerate(rows)
                ],
                colWidths=widths,
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9ECEF")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C7D4DA")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story += [table, Spacer(1, 10)]
            word_table = document.add_table(rows=0, cols=len(rows[0]))
            word_table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                cells = word_table.add_row().cells
                for cell, t in zip(cells, row, strict=True):
                    word_inline(cell.paragraphs[0], t)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = DocPt(8)
                        if row_index == 0:
                            run.bold = True
                # Evita righe spezzate e ripete l'intestazione nelle pagine successive.
                props = word_table.rows[-1]._tr.get_or_add_trPr()
                props.append(
                    etree.Element(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit"
                    )
                )
                if row_index == 0:
                    props.append(
                        etree.Element(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader"
                        )
                    )
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            lists.append(
                {
                    "ordered": token.type == "ordered_list_open",
                    "number": int(token.attrGet("start") or 1),
                }
            )
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif token.type == "paragraph_open":
            content = tokens[i + 1]
            prefix = ""
            if lists:
                prefix = f"{lists[-1]['number']}. " if lists[-1]["ordered"] else "• "
                lists[-1]["number"] += 1
            story.append(Paragraph(prefix + inline(content), styles["Normal"]))
            word_inline(document.add_paragraph(), content, prefix)
            i += 3
            continue
        elif token.type == "fence":
            if token.info == "mermaid":
                story.append(Image(str(HERE / "architettura.png"), width=475, height=213.75))
                document.add_picture(str(HERE / "architettura.png"), width=DocInches(6.3))
            else:
                story.append(
                    KeepTogether([XPreformatted(html.escape(token.content.rstrip()), code_style)])
                )
                p = document.add_paragraph()
                p.paragraph_format.keep_together = True
                run = p.add_run(token.content.rstrip())
                run.font.name = "Courier New"
                run.font.size = DocPt(8)
        i += 1

    # Mantieni titolo o frase introduttiva insieme al blocco di codice successivo.
    grouped_story = []
    for flowable in story:
        if (
            isinstance(flowable, KeepTogether)
            and grouped_story
            and isinstance(grouped_story[-1], Paragraph)
        ):
            flowable = KeepTogether([grouped_story.pop(), *flowable._content])
        grouped_story.append(flowable)

    pdf = SimpleDocTemplate(
        str(HERE / "relazione.pdf"),
        pagesize=A4,
        leftMargin=60,
        rightMargin=60,
        topMargin=48,
        bottomMargin=48,
        title=title,
        author=cover["autore"],
    )
    pdf.build(grouped_story)
    set_theme_fonts(document.part.package)
    document.save(HERE / "relazione.docx")


def create_viewer():
    cases = {}
    for folder in sorted((HERE / "evidenze").glob("demo_*-r1")):
        cases[folder.name] = {
            p.stem: json.loads(p.read_text()) for p in sorted(folder.glob("*.json"))
        }
    payload = json.dumps(cases, ensure_ascii=False).replace("<", "\\u003c")
    template = (HERE / "viewer.template.html").read_text()
    (HERE / "demo-evidenze.html").write_text(template.replace("PAYLOAD", payload))


def create_notes():
    lines = [
        "# Copione delle slide",
        "",
        "14 minuti + 1 minuto per domande. Le appendici sono facoltative.",
        "",
    ]
    elapsed = 0
    for n, slide in enumerate(SLIDES, 1):
        end = elapsed + slide["seconds"]
        stamp = f"{elapsed // 60:02d}:{elapsed % 60:02d}–{end // 60:02d}:{end % 60:02d}"
        lines += [
            f"## {n}. {slide['title']}",
            "",
            stamp if slide["seconds"] else "Appendice",
            "",
            slide["notes"],
            "",
        ]
        elapsed = end
    assert elapsed == 840
    (HERE / "copione-slide.md").write_text("\n".join(lines))


if __name__ == "__main__":
    create_pptx()
    font = register_fonts()
    create_figure(font)
    create_report(font)
    create_viewer()
    create_notes()
    print("Creati PowerPoint, relazione PDF/DOCX, figura, viewer offline e copione.")
